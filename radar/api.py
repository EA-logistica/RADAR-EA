import csv
import io
import re
import secrets
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Annotated, Literal
from fastapi import FastAPI, Depends, HTTPException, Query, Header
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field, model_validator
from sqlalchemy import select, func, or_, case, distinct, cast, String
from sqlalchemy.orm import Session as DBSession
from openpyxl import Workbook
from radar.config import settings
from radar.db import session_dependency
from radar.models import Operation as O, Artifact, Run, Review, Revision, Alert, now
from radar.normalize import text, search_groups, MATERIALS
from radar.grades import grade_key, APPLICATIONS, parse_product_query
from radar.etl import create_run

app=FastAPI(title='Radar de Importaciones',version='0.1.0',description='Series aduaneras con trazabilidad. FOB y CIF declarados, no costo puesto en almacén.')
DB=Annotated[DBSession,Depends(session_dependency)]

def authorize(x_radar_token: str | None=Header(default=None)):
    if settings.api_token and not secrets.compare_digest(x_radar_token or '',settings.api_token):
        raise HTTPException(401,'Token de acceso requerido')

class Filters(BaseModel):
    q: str = Field(default='',max_length=200)
    importer: str = ''
    ruc: str = ''
    supplier: str = ''
    origin: str = ''
    material: str = ''
    hs: str = ''
    start: date | None = None
    end: date | None = None
    review: bool = False
    application: str = ''
    brand: str = ''
    grade: str = ''
    mi_min: Decimal | None = None
    mi_max: Decimal | None = None
    scope: Literal['plastics','all'] = 'plastics'

def conditions(f):
    c=[]
    if f.scope=='plastics': c.append(O.plastics_scope.is_(True))
    if f.start: c.append(O.numbered_on>=f.start)
    if f.end: c.append(O.numbered_on<=f.end)
    if f.start and f.end and f.start>f.end: raise HTTPException(422,'Rango de fechas inválido')
    for field,value in [(O.importer,text(f.importer)),(O.supplier,text(f.supplier))]:
        if value: c.append(field.contains(value,autoescape=True))
    for field,value in [(O.importer_ruc,f.ruc),(O.origin,f.origin),(O.material,f.material)]:
        if value: c.append(field==value)
    if f.hs: c.append(O.hs_code.startswith(f.hs,autoescape=True))
    if f.review: c.append(O.needs_review.is_(True))
    if f.application: c.append(O.applications.contains('|'+f.application+'|',autoescape=True))
    if f.brand: c.append(O.brand.ilike('%'+f.brand.replace('%','').replace('_','')+'%'))
    if f.grade: c.append(O.grade_key==grade_key(f.grade))
    if f.mi_min is not None: c.append(O.melt_index>=f.mi_min)
    if f.mi_max is not None: c.append(O.melt_index<=f.mi_max)
    for group in search_groups(f.q):
        # PostgreSQL word boundaries prevent PP in unrelated terms and 0.35 matching 0.3509.
        searchable=O.search_text+' '+func.coalesce(O.grade_text,'')
        c.append(or_(*(searchable.op('~')(r'\m'+re.escape(term)+r'\M') for term in group)))
    return c

def operation_dict(row, full=False):
    fields=['id','customs','year','declaration','series','numbered_on','importer_ruc','importer','supplier','supplier_status','origin','acquisition_country','hs_code','description','material','classification_reason','needs_review','classification_locked','currency','quantity','unit','net_kg','kg_method','fob_usd','freight_usd','insurance_usd','cif_usd','usd_kg','brand','grade','grade_key','application','applications','melt_index','density','product_name','grade_info','quality_flags','source_url','artifact_id','updated_at']
    data={k:getattr(row,k) for k in fields}
    if full: data['raw']=row.raw
    return data

def declaration_key(): return O.customs+cast(O.year,String)+O.regime+O.declaration

def summary(s,c):
    paired=O.usd_kg.is_not(None)
    result=s.execute(select(func.count(O.id),func.count(distinct(declaration_key())),func.sum(O.fob_usd),func.sum(O.net_kg),func.count(distinct(O.importer_ruc)),func.count(distinct(O.supplier)),func.count(case((O.needs_review,1))),func.count(case((O.supplier.is_(None),1))),func.sum(case((paired,O.fob_usd))),func.sum(case((paired,O.net_kg))),func.min(O.numbered_on),func.max(O.numbered_on)).where(*c)).one()
    return dict(series=result[0],operations=result[1],fob_usd=result[2],tonnes=result[3]/1000 if result[3] is not None else None,importers=result[4],suppliers=result[5],pending_review=result[6],missing_supplier=result[7],usd_kg=result[8]/result[9] if result[9] else None,start=result[10],end=result[11])

def trend(s,c):
    day=func.date_trunc('week',O.numbered_on)
    paired=O.usd_kg.is_not(None)
    rows=s.execute(select(day.label('period'),func.sum(O.fob_usd).label('fob_usd'),(func.sum(O.net_kg)/1000).label('tonnes'),(func.sum(case((paired,O.fob_usd)))/func.nullif(func.sum(case((paired,O.net_kg))),0)).label('usd_kg'),func.count(O.id).label('series')).where(*c).group_by(day).order_by(day)).mappings().all()
    return [dict(r) for r in rows]

def breakdown(s,c,field,limit=10):
    rows=s.execute(select(field.label('name'),func.sum(O.fob_usd).label('fob_usd'),(func.sum(O.net_kg)/1000).label('tonnes'),func.count(O.id).label('series')).where(*c).group_by(field).order_by(func.sum(O.fob_usd).desc().nullslast()).limit(limit)).mappings().all()
    return [dict(r) for r in rows]

def cost_columns():
    # CFR = FOB + flete; CIF = CFR + seguro. Valores declarados, no costo puesto en almacén.
    paired=O.usd_kg.is_not(None)
    cfr=O.fob_usd+O.freight_usd
    kg=func.nullif(func.sum(case((paired,O.net_kg))),0)
    return [func.count(O.id).label('series'),(func.sum(O.net_kg)/1000).label('tonnes'),func.sum(O.fob_usd).label('fob_usd'),func.sum(O.freight_usd).label('freight_usd'),func.sum(O.insurance_usd).label('insurance_usd'),func.sum(cfr).label('cfr_usd'),func.sum(O.cif_usd).label('cif_usd'),
        (func.sum(case((paired,O.fob_usd)))/kg).label('fob_kg'),(func.sum(case((paired,cfr)))/kg).label('cfr_kg'),(func.sum(case((paired,O.cif_usd)))/kg).label('cif_kg'),
        (100*func.sum(O.freight_usd)/func.nullif(func.sum(O.fob_usd),0)).label('freight_pct'),(100*func.sum(O.insurance_usd)/func.nullif(func.sum(O.fob_usd),0)).label('insurance_pct')]

def cost_summary(s,c):
    return dict(s.execute(select(*cost_columns()).where(*c)).mappings().one())

auth=[Depends(authorize)]

@app.get('/api/health')
def health(s:DB):
    s.execute(select(1));return {'status':'ok','database':'postgresql'}

@app.get('/api/dashboard',dependencies=auth)
def dashboard(s:DB,f:Annotated[Filters,Depends()]):
    c=conditions(f)
    return {'summary':summary(s,c),'trend':trend(s,c),'materials':breakdown(s,c,O.material),'countries':breakdown(s,c,O.origin),'importers':breakdown(s,c,O.importer),'suppliers':breakdown(s,c,O.supplier),'costs':cost_summary(s,c)}

@app.get('/api/costs',dependencies=auth)
def costs(s:DB,f:Annotated[Filters,Depends()],group:Literal['declaration','material','origin','importer']='declaration',sort:Literal['cif_usd','cif_kg','freight_pct','numbered_on']='cif_usd',page:int=Query(1,ge=1),page_size:int=Query(15,ge=1,le=100)):
    c=conditions(f)
    if group=='declaration':
        keys=[O.customs,O.year,O.regime,O.declaration]
        extra=[O.customs,O.year,O.declaration,func.min(O.numbered_on).label('numbered_on'),func.max(O.importer_ruc).label('importer_ruc'),func.max(O.importer).label('importer'),func.string_agg(distinct(O.origin),', ').label('origin'),func.string_agg(distinct(O.material),', ').label('material')]
    elif group=='importer':
        keys=[O.importer_ruc];extra=[O.importer_ruc,func.max(O.importer).label('importer')]
    else:
        field=O.material if group=='material' else O.origin
        keys=[field];extra=[field.label('name')]
    columns=cost_columns()
    order={col.name:col for col in columns}.get(sort)
    if sort=='numbered_on' and group=='declaration': order=extra[3]
    if order is None: order=columns[6]
    total=s.scalar(select(func.count()).select_from(select(*keys).where(*c).group_by(*keys).subquery()))
    rows=s.execute(select(*extra,*columns).where(*c).group_by(*keys).order_by(order.desc().nullslast(),*keys).offset((page-1)*page_size).limit(page_size)).mappings().all()
    return {'total':total,'page':page,'page_size':page_size,'group':group,'items':[dict(r) for r in rows]}

def grade_rows(s,c,sort='cif_usd',page=1,page_size=20,importer_detail=False,prefer_app=None):
    """Un registro por grado (código de materia prima) con su ficha, descripción corta y quiénes lo importan."""
    c=[*c,O.grade_key.is_not(None)]
    paired=O.usd_kg.is_not(None)
    mode=lambda col: func.mode().within_group(col)
    cols=[O.grade_key.label('grade_key'),mode(O.grade).label('grade'),mode(O.brand).label('brand'),mode(O.product_name).label('product_name'),mode(O.material).label('material'),
        mode(O.application).label('application'),func.string_agg(distinct(O.applications),'').label('applications'),
        func.min(O.melt_index).label('mi_min'),func.max(O.melt_index).label('mi_max'),mode(O.melt_index).label('melt_index'),func.max(O.density).label('density'),
        func.count(O.id).label('series'),func.count(distinct(declaration_key())).label('operations'),func.count(distinct(O.importer_ruc)).label('importers'),
        func.string_agg(distinct(O.importer),' · ').label('importer_names'),(func.sum(O.net_kg)/1000).label('tonnes'),func.sum(O.fob_usd).label('fob_usd'),func.sum(O.cif_usd).label('cif_usd'),
        (func.sum(case((paired,O.cif_usd)))/func.nullif(func.sum(case((paired,O.net_kg))),0)).label('cif_kg'),func.min(O.numbered_on).label('first'),func.max(O.numbered_on).label('last')]
    order={'cif_usd':cols[17],'tonnes':cols[15],'importers':cols[13],'cif_kg':cols[18],'melt_index':cols[9],'last':cols[20]}[sort]
    total=s.scalar(select(func.count(distinct(O.grade_key))).where(*c))
    rows=s.execute(select(*cols).where(*c).group_by(O.grade_key).order_by(*([case((cols[5]==prefer_app,0),else_=1)] if prefer_app else []),order.desc().nullslast(),O.grade_key).offset((page-1)*page_size).limit(page_size)).mappings().all()
    keys=[r['grade_key'] for r in rows]
    info,buyers={},{}
    if keys:
        # Ficha más informativa del grado: catálogo > MI conocido > primera serie.
        rank=lambda gi:(bool(gi.get('catalog')),gi.get('mi_source') is not None,bool(gi.get('summary')))
        for k,gi in s.execute(select(O.grade_key,O.grade_info).where(O.grade_key.in_(keys)).order_by(O.grade_key,O.id)):
            gi=gi or {}
            if k not in info or rank(gi)>rank(info[k]): info[k]=gi
        if importer_detail:
            for r in s.execute(select(O.grade_key,O.importer_ruc,func.max(O.importer).label('importer'),func.count(O.id).label('series'),(func.sum(O.net_kg)/1000).label('tonnes'),func.sum(O.cif_usd).label('cif_usd'),
                    (func.sum(case((paired,O.cif_usd)))/func.nullif(func.sum(case((paired,O.net_kg))),0)).label('cif_kg'),func.max(O.numbered_on).label('last'),func.string_agg(distinct(O.origin),', ').label('origins'))
                    .where(*c,O.grade_key.in_(keys)).group_by(O.grade_key,O.importer_ruc).order_by(O.grade_key,func.sum(O.net_kg).desc().nullslast())).mappings():
                buyers.setdefault(r['grade_key'],[]).append({k:v for k,v in r.items() if k!='grade_key'})
    items=[]
    for r in rows:
        d=dict(r);gi=info.get(r['grade_key']) or {}
        d['applications']=[a for a in dict.fromkeys((d['applications'] or '').split('|')) if a]
        d['importer_names']=(d['importer_names'] or '').split(' · ')
        d.update({k:gi.get(k) for k in ('polymer','family','summary','mi_hint','mi_condition','mi_source','application_source','source_url','confidence','catalog','catalog_applications','iv','application_mismatch')})
        if importer_detail: d['buyers']=buyers.get(r['grade_key'],[])
        items.append(d)
    return {'total':total,'page':page,'page_size':page_size,'items':items}

@app.get('/api/grades',dependencies=auth)
def grades(s:DB,f:Annotated[Filters,Depends()],sort:Literal['cif_usd','tonnes','importers','cif_kg','melt_index','last']='cif_usd',page:int=Query(1,ge=1),page_size:int=Query(20,ge=1,le=100)):
    return grade_rows(s,conditions(f),sort,page,page_size)

def product_conditions(p):
    """Filtros estructurados del buscador de productos: familia, proceso, MI y texto libre (marca o grado)."""
    c=[]
    if p['family']: c.append(or_(O.material==p['family'],O.grade_info['family'].as_string()==p['family']))
    if p['application']: c.append(O.applications.contains('|'+p['application']+'|',autoescape=True))
    if p['mi_min'] is not None: c.append(O.melt_index>=p['mi_min'])
    if p['mi_max'] is not None: c.append(O.melt_index<=p['mi_max'])
    searchable=O.search_text+' '+func.coalesce(O.grade_text,'')
    for term in p['terms']:
        c.append(or_(searchable.op('~')(r'\m'+re.escape(term)+r'\M'),O.grade_key==grade_key(term)))
    return c

@app.get('/api/products',dependencies=auth)
def products(s:DB,f:Annotated[Filters,Depends()],p:str=Query('',max_length=200),sort:Literal['cif_usd','tonnes','importers','cif_kg','melt_index','last']='tonnes',page:int=Query(1,ge=1),page_size:int=Query(12,ge=1,le=50)):
    """Buscador de productos: 'hdpe inyeccion' -> grados que cumplen, con descripción y MI, y todas las empresas que los importaron."""
    parsed=parse_product_query(p)
    c=[*conditions(f),*product_conditions(parsed)]
    paired=O.usd_kg.is_not(None)
    cif_kg=(func.sum(case((paired,O.cif_usd)))/func.nullif(func.sum(case((paired,O.net_kg))),0))
    totals=s.execute(select(func.count(O.id),func.count(distinct(declaration_key())),func.count(distinct(O.importer_ruc)),func.sum(O.net_kg)/1000,func.sum(O.cif_usd),cif_kg,
        func.count(distinct(O.grade_key)),func.count(case((O.grade_key.is_(None),1))),func.min(O.numbered_on),func.max(O.numbered_on)).where(*c)).one()
    summary=dict(zip(['series','operations','importers','tonnes','cif_usd','cif_kg','grades','ungraded_series','start','end'],totals))
    importers=[dict(r) for r in s.execute(select(O.importer_ruc.label('ruc'),func.max(O.importer).label('name'),func.count(O.id).label('series'),func.count(distinct(declaration_key())).label('operations'),
        (func.sum(O.net_kg)/1000).label('tonnes'),func.sum(O.cif_usd).label('cif_usd'),cif_kg.label('cif_kg'),func.max(O.numbered_on).label('last'),
        func.string_agg(distinct(O.grade),' · ').label('grades'),func.string_agg(distinct(O.brand),' · ').label('brands'),func.string_agg(distinct(O.origin),', ').label('origins'))
        .where(*c).group_by(O.importer_ruc).order_by(func.sum(O.net_kg).desc().nullslast()).limit(200)).mappings()]
    brands=[dict(r) for r in s.execute(select(O.brand.label('name'),func.count(O.id).label('series'),(func.sum(O.net_kg)/1000).label('tonnes'),func.sum(O.cif_usd).label('cif_usd'),
        func.count(distinct(O.grade_key)).label('grades')).where(*c,O.brand.is_not(None)).group_by(O.brand).order_by(func.sum(O.net_kg).desc().nullslast()).limit(15)).mappings()]
    return {'query':p,'interpretation':parsed,'summary':summary,'importers':importers,'brands':brands,'grades':grade_rows(s,c,sort,page,page_size,importer_detail=True,prefer_app=parsed['application'])}

@app.get('/api/operations',dependencies=auth)
def operations(s:DB,f:Annotated[Filters,Depends()],page:int=Query(1,ge=1),page_size:int=Query(25,ge=1,le=100)):
    c=conditions(f);total=s.scalar(select(func.count()).select_from(O).where(*c))
    rows=s.scalars(select(O).where(*c).order_by(O.numbered_on.desc(),O.id.desc()).offset((page-1)*page_size).limit(page_size))
    return {'total':total,'page':page,'page_size':page_size,'items':[operation_dict(r) for r in rows]}

@app.get('/api/analytics',dependencies=auth)
def analytics(s:DB,f:Annotated[Filters,Depends()],grain:Literal['day','week','month']='week',compare_start:date | None=None,compare_end:date | None=None):
    from radar.analytics import history
    return history(s,f,conditions,grain,compare_start,compare_end)

@app.get('/api/analytics/export',dependencies=auth)
def analytics_export(s:DB,f:Annotated[Filters,Depends()],grain:Literal['day','week','month']='week'):
    from radar.analytics import history
    result=history(s,f,conditions,grain)
    out=io.StringIO();writer=csv.writer(out)
    fields=['period','end','selection_start','selection_end','status','fob_usd','tonnes','operations','series','importers','usd_kg','observed_cumulative_fob','fob_change_percent']
    writer.writerow(fields+['days','backed_days','coverage_note'])
    for row in result['timeline']:
        writer.writerow([safe_cell(row.get(k)) for k in fields]+[row['coverage']['days'],row['coverage']['backed_days'],result['meta']['coverage_note']])
    return StreamingResponse(iter([out.getvalue().encode('utf-8-sig')]),media_type='text/csv; charset=utf-8',headers={'Content-Disposition':'attachment; filename="radar-historico.csv"'})

@app.get('/api/operations/{op_id}',dependencies=auth)
def operation_detail(op_id:int,s:DB):
    row=s.get(O,op_id)
    if not row: raise HTTPException(404,'Serie no encontrada')
    data=operation_dict(row,True)
    data['artifact']=s.get(Artifact,row.artifact_id)
    data['revisions']=list(s.scalars(select(Revision).where(Revision.operation_id==op_id).order_by(Revision.observed_at.desc())))
    data['reviews']=list(s.scalars(select(Review).where(Review.operation_id==op_id).order_by(Review.created_at.desc())))
    return data

@app.get('/api/options',dependencies=auth)
def options(s:DB):
    return {'applications':[a for a,_ in APPLICATIONS],'brands':list(s.scalars(select(O.brand).where(O.brand.is_not(None),O.plastics_scope.is_(True)).group_by(O.brand).having(func.count()>=2).order_by(O.brand))),'materials':MATERIALS,'origins':list(s.scalars(select(O.origin).where(O.origin.is_not(None)).distinct().order_by(O.origin))),'suppliers':list(s.scalars(select(O.supplier).where(O.supplier.is_not(None)).distinct().order_by(O.supplier)))}

@app.get('/api/companies',dependencies=auth)
def companies(s:DB,f:Annotated[Filters,Depends()]):
    rows=s.execute(select(O.importer_ruc.label('ruc'),func.max(O.importer).label('name'),func.sum(O.fob_usd).label('fob_usd'),(func.sum(O.net_kg)/1000).label('tonnes'),func.count(distinct(declaration_key())).label('operations'),func.min(O.numbered_on).label('first'),func.max(O.numbered_on).label('last')).where(*conditions(f),O.importer_ruc.is_not(None)).group_by(O.importer_ruc).order_by(func.sum(O.fob_usd).desc().nullslast()).limit(500)).mappings().all()
    return [dict(r) for r in rows]

@app.get('/api/company/{ruc}',dependencies=auth)
def company(ruc:str,s:DB,f:Annotated[Filters,Depends()]):
    c=conditions(f)+[O.importer_ruc==ruc]
    dates=list(s.scalars(select(O.numbered_on).where(*c).distinct().order_by(O.numbered_on)))
    gaps=[(b-a).days for a,b in zip(dates,dates[1:])]
    return {'name':s.scalar(select(func.max(O.importer)).where(*c)),'ruc':ruc,'summary':summary(s,c),'trend':trend(s,c),'materials':breakdown(s,c,O.material),'suppliers':breakdown(s,c,O.supplier),'countries':breakdown(s,c,O.origin),'mean_days_between_active_dates':sum(gaps)/len(gaps) if gaps else None}

@app.get('/api/material/{material}',dependencies=auth)
def material_profile(material:str,s:DB,f:Annotated[Filters,Depends()]):
    c=conditions(f)+[O.material==material]
    return {'name':material,'summary':summary(s,c),'trend':trend(s,c),'importers':breakdown(s,c,O.importer),'suppliers':breakdown(s,c,O.supplier),'countries':breakdown(s,c,O.origin)}

@app.get('/api/compare',dependencies=auth)
def compare(s:DB,f:Annotated[Filters,Depends()],rucs:str=''):
    selected=list(dict.fromkeys(rucs.split(',')))
    if not 2<=len(selected)<=4 or any(not re.fullmatch(r'\d{11}',r) for r in selected): raise HTTPException(422,'Selecciona de 2 a 4 RUC')
    return [company(ruc,s,f) for ruc in selected]

class ReviewInput(BaseModel):
    material: str
    note: str=Field(min_length=5,max_length=1000)
    actor: str=Field(default='Analista local',min_length=2,max_length=100)

@app.post('/api/operations/{op_id}/review',dependencies=auth)
def review(op_id:int,body:ReviewInput,s:DB):
    if body.material not in MATERIALS: raise HTTPException(422,'Material no permitido')
    op=s.get(O,op_id)
    if not op: raise HTTPException(404,'Serie no encontrada')
    s.add(Review(operation_id=op.id,previous_material=op.material,material=body.material,note=body.note,actor=body.actor))
    op.material=body.material;op.classification_locked=True;op.needs_review=body.material=='Revisar';op.classification_reason='Revisión humana: '+body.note;op.updated_at=now();s.commit()
    return {'status':'saved'}

def safe_cell(value):
    if value is None: return ''
    if isinstance(value,str) and value.lstrip().startswith(('=','+','-','@')): return "'"+value
    return value

@app.get('/api/export',dependencies=auth)
def export(s:DB,f:Annotated[Filters,Depends()],format:Literal['csv','xlsx']='csv'):
    c=conditions(f)
    count=s.scalar(select(func.count()).select_from(O).where(*c))
    if count>100000: raise HTTPException(422,'Reduce filtros a 100 000 series para exportar')
    columns=['numbered_on','customs','year','declaration','series','importer_ruc','importer','supplier','supplier_status','origin','material','hs_code','description','net_kg','fob_usd','freight_usd','insurance_usd','cif_usd','usd_kg','kg_method','source_url','artifact_id']
    rows=s.scalars(select(O).where(*c).order_by(O.numbered_on,O.id))
    headers={'Content-Disposition':f'attachment; filename="radar-importaciones.{format}"'}
    if format=='csv':
        out=io.StringIO();writer=csv.writer(out);writer.writerow(columns)
        for row in rows: writer.writerow([safe_cell(getattr(row,k)) for k in columns])
        return StreamingResponse(iter([out.getvalue().encode('utf-8-sig')]),media_type='text/csv; charset=utf-8',headers=headers)
    wb=Workbook(write_only=True);ws=wb.create_sheet('Series');ws.append(columns)
    for row in rows: ws.append([safe_cell(getattr(row,k)) for k in columns])
    meta=wb.create_sheet('Metodología');meta.append(['FOB USD/kg','Valor declarado de la misma serie / kg verificables. No es costo en almacén.']);meta.append(['Cobertura','Sólo archivos y consultas cargados. Proveedor vacío no significa inexistente.'])
    out=io.BytesIO();wb.save(out);out.seek(0)
    return StreamingResponse(out,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers=headers)

@app.get('/api/quality',dependencies=auth)
def quality(s:DB):
    coverage=summary(s,[O.plastics_scope.is_(True)])
    last=s.scalar(select(func.max(Artifact.downloaded_at)))
    files=list(s.scalars(select(Artifact).order_by(Artifact.downloaded_at.desc())))
    runs=list(s.scalars(select(Run).order_by(Run.started_at.desc()).limit(30)))
    return {'coverage':coverage,'last_download':last,'artifacts':files,'runs':runs,'scope':'Importación definitiva, series y períodos efectivamente cargados; no cobertura nacional completa.','freshness':'Bases: publicación semanal anunciada por SUNAT. Consultas: frecuencia no verificada.','supplier_note':'SUNAT puede ocultar proveedores. No se infieren a partir de marcas.','price_note':'FOB USD/kg ponderado por kg de series comparables. No es precio final en almacén.'}

@app.get('/api/artifacts/{artifact_id}/download',dependencies=auth)
def artifact_download(artifact_id:str,s:DB):
    item=s.get(Artifact,artifact_id)
    if not item or not Path(item.path).is_file(): raise HTTPException(404,'Original no disponible en este entorno')
    return FileResponse(item.path,filename=Path(item.path).name,media_type='application/octet-stream')

@app.get('/api/alerts',dependencies=auth)
def alerts(s:DB):
    return list(s.scalars(select(Alert).order_by(Alert.created_at.desc()).limit(100)))

@app.post('/api/alerts/{alert_id}/read',dependencies=auth)
def mark_alert(alert_id:int,s:DB):
    item=s.get(Alert,alert_id)
    if not item: raise HTTPException(404,'Alerta no encontrada')
    item.read=True;s.commit();return {'status':'saved'}

class RunInput(BaseModel):
    kind: Literal['bulk','query']='bulk'
    weeks:int=Field(default=1,ge=1,le=12)
    query_kind:Literal['importer','hs']='importer'
    value:str=''
    start:date | None=None
    end:date | None=None
    force:bool=False
    scope:Literal['plastics','all']='plastics'
    @model_validator(mode='after')
    def validate_query(self):
        if self.kind=='query':
            length=11 if self.query_kind=='importer' else 10
            if not re.fullmatch(r'\d{'+str(length)+'}',self.value): raise ValueError('RUC/subpartida inválido')
            if not self.start or not self.end or self.end<self.start or (self.end-self.start).days>366: raise ValueError('Fechas de consulta inválidas')
        return self

@app.post('/api/runs',dependencies=auth,status_code=202)
def enqueue(body:RunInput,s:DB):
    if s.scalar(select(func.count()).select_from(Run).where(Run.status.in_(['queued','running'])))>=5: raise HTTPException(409,'Espera a que finalicen las cargas pendientes')
    params={'weeks':body.weeks,'scope':body.scope,'force':body.force} if body.kind=='bulk' else {'kind':body.query_kind,'value':body.value,'start':body.start.isoformat(),'end':body.end.isoformat(),'force':body.force}
    return {'id':create_run(body.kind,params),'status':'queued'}

@app.post('/api/runs/{run_id}/resume',dependencies=auth)
def resume(run_id:str,s:DB):
    run=s.get(Run,run_id)
    if not run: raise HTTPException(404,'Ejecución no encontrada')
    if run.status!='failed': raise HTTPException(409,'Sólo se reanudan ejecuciones fallidas')
    run.status='queued';run.error=None;s.commit();return {'status':'queued'}

if settings.web_dist.exists():
    app.mount('/',StaticFiles(directory=settings.web_dist,html=True),name='web')

